"""Text extraction for supported document formats.

Extraction is intentionally conservative: no macros, no embedded objects, no
external references are followed. Unsupported or malformed files fail with a
safe error code rather than propagating parser internals.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger("apps.knowledge.extraction")


class ExtractionError(Exception):
    """Raised when a document cannot be parsed; message is safe for audit."""

    def __init__(self, message: str, *, code: str = "extraction_failed"):
        super().__init__(message)
        self.code = code


@dataclass
class TextSegment:
    text: str
    page_number: int | None = None


@dataclass
class ExtractedDocument:
    segments: list[TextSegment] = field(default_factory=list)
    page_count: int = 0

    @property
    def text(self) -> str:
        return "\n\n".join(segment.text for segment in self.segments if segment.text.strip())


def _extract_pdf(path: Path) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ExtractionError("Encrypted PDF documents are not supported.", code="encrypted_document")
        segments: list[TextSegment] = []
        for index, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001 - one bad page must not fail the document
                logger.warning("Failed to extract page %s", index)
                text = ""
            if text.strip():
                segments.append(TextSegment(text=text, page_number=index))
        return ExtractedDocument(segments=segments, page_count=len(reader.pages))
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The PDF could not be parsed.", code="pdf_parse_error") from exc


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise ExtractionError("DOCX support is not installed.", code="unsupported_format") from exc

    try:
        document = docx.Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        # Tables often carry the useful content in business documents.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return ExtractedDocument(segments=[TextSegment(text="\n\n".join(paragraphs))], page_count=0)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The DOCX file could not be parsed.", code="docx_parse_error") from exc


def _extract_html(path: Path) -> ExtractedDocument:
    from bs4 import BeautifulSoup

    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript", "iframe", "object", "embed"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines()]
        cleaned = "\n".join(line for line in lines if line)
        return ExtractedDocument(segments=[TextSegment(text=cleaned)])
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The HTML file could not be parsed.", code="html_parse_error") from exc


def _extract_plain(path: Path) -> ExtractedDocument:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The text file could not be read.", code="text_parse_error") from exc
    return ExtractedDocument(segments=[TextSegment(text=text)])


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".html": _extract_html,
    ".htm": _extract_html,
    ".txt": _extract_plain,
    ".md": _extract_plain,
    ".markdown": _extract_plain,
}


def extract_document(path: str | Path) -> ExtractedDocument:
    file_path = Path(path)
    extension = file_path.suffix.lower()
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise ExtractionError("This file type is not supported.", code="unsupported_format")

    result = extractor(file_path)
    if not result.text.strip():
        raise ExtractionError("No readable text was found in the document.", code="empty_document")
    return result
